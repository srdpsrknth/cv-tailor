from pathlib import Path

from docx import Document

from models.profile import Profile
from models.tailored_output import TailoredCoverLetter
from renderers.docx_renderer import DocxRenderer


def test_cover_letter_preserves_xml_sensitive_text(tmp_path, monkeypatch):
    renderer = DocxRenderer(outputs_dir=str(tmp_path))
    monkeypatch.setattr(renderer, '_generate_pdf_via_latex', lambda *args: tmp_path / 'letter.pdf')
    profile = Profile(version='1.0', identity={'full_name': 'Candidate', 'email': 'a@example.com'})
    letter = TailoredCoverLetter(company_name='Data & Capital', role_title='Engineer',
        paragraphs=['Data & Technology, where 1 < 2 and 3 > 2.', 'Second.', 'Third.', 'Fourth.'],
        signature_name='Candidate')
    path, _ = renderer.render_cover_letter(letter, profile)
    text = '\n'.join(p.text for p in Document(path).paragraphs)
    assert 'Data & Capital' in text
    assert letter.paragraphs[0] in text
